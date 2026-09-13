from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm
from fastapi.testclient import TestClient

from app.catalog import check_contiguous_placements, document_types, templates
from app.main import create_app
from app.schemas import DocumentType
from app.settings import Settings

DRAFT = (
    "Иванов И. И. просит согласовать 30 000 рублей до 25.09.2026.\n"
    "Поставка возможна только после согласования бюджета."
)
REQUISITES = {
    "organization": "ФГБОУ ВО «БГИТУ»",
    "recipient": "Руководителю учебного отдела Сидоровой А. В.",
    "sender": "Заведующего лабораторией Иванова И. И.",
    "subject": "О закупке мониторов",
    "date": "11.09.2026",
    "number": "47-СЗ",
    "salutation": "Уважаемая Анна Викторовна!",
    "attachment": "акт на 1 л. в 1 экз.",
    "purpose": "для планово-финансового отдела",
    "signer_position": "Заведующий лабораторией",
    "signer": "И. И. Иванов",
    "executor": "Петров П. П., (4832) 00-00-01",
}


@pytest.fixture
def client():
    with TestClient(create_app(Settings(text_processor="stub"))) as client:
        yield client


def prepared(client, type_id="service_memo", requisites=None, draft=DRAFT):
    response = client.post("/api/process", json={
        "draft": draft, "doc_type": type_id, "requisites": requisites or {},
    })
    assert response.status_code == 200, response.text
    return response.json()


def download(client, document, template_id="classic", formatting=None):
    payload = {"document": document, "template_id": template_id}
    if formatting is not None:
        payload["formatting"] = formatting
    response = client.post("/api/documents/download", json=payload)
    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].endswith('.docx"')
    assert "wordprocessingml.document" in response.headers["content-type"]
    return response


def filled(type_id):
    return {field.id: REQUISITES[field.id] for field in document_types()[type_id].fields}


def docx(client, type_id="service_memo", template_id="classic", requisites=None):
    result = prepared(client, type_id, requisites)
    return Document(BytesIO(download(client, result["document"], template_id).content))


def texts_of(doc) -> list[str]:
    return [paragraph.text for paragraph in doc.paragraphs]


def document_text(content: bytes) -> str:
    """Paragraphs, table cells, page headers and footers."""
    doc = Document(BytesIO(content))
    parts = texts_of(doc)
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts += texts_of(part)
    return "\n".join(parts)


def test_catalog_and_health(client):
    catalog = client.get("/api/catalog").json()
    assert len(catalog["doc_types"]) == 4
    assert len(catalog["templates"]) == 2
    assert client.get("/api/health").json()["status"] == "ok"


@pytest.mark.parametrize("type_id", list(document_types()))
@pytest.mark.parametrize("template_id", list(templates()))
def test_all_types_and_templates_generate_editable_documents(client, type_id, template_id):
    result = prepared(client, type_id)
    response = download(client, result["document"], template_id)
    text = document_text(response.content)
    assert document_types()[type_id].title in text
    for line in DRAFT.splitlines():
        assert line in text
    for field in result["missing_fields"]:
        assert f'[Заполнить: {field["label"]}]' in text
    with ZipFile(BytesIO(response.content)) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())


def test_stub_preserves_facts_and_does_not_claim_corrections(client):
    result = prepared(client)
    assert result["document"]["body"] == DRAFT.splitlines()
    assert result["document"]["requisites"] == {}
    assert result["changes"] == []
    assert result["processor_mode"] == "stub"
    assert result["warnings"]


@pytest.mark.parametrize("template_id", list(templates()))
def test_supplied_requisites_replace_placeholders(client, template_id):
    fields = document_types()["service_memo"].fields
    values = {field.id: f"Введено пользователем: {field.label}" for field in fields}
    result = prepared(client, requisites=values)
    assert not result["missing_fields"]
    text = document_text(download(client, result["document"], template_id).content)
    assert "[Заполнить:" not in text
    for value in values.values():
        assert value in text


def test_many_short_lines_within_draft_limit_are_valid(client):
    response = client.post("/api/process", json={
        "draft": "\n".join(["x"] * 1001), "doc_type": "service_memo", "requisites": {},
    })
    assert response.status_code == 200
    assert len(response.json()["document"]["body"]) == 1001


def test_switching_template_changes_formatting_only(client):
    result = prepared(client, requisites=filled("service_memo"))
    classic_bytes = download(client, result["document"], "classic").content
    modern_bytes = download(client, result["document"], "modern").content
    classic, modern = Document(BytesIO(classic_bytes)), Document(BytesIO(modern_bytes))
    for value in [*result["document"]["requisites"].values(), *DRAFT.splitlines()]:
        assert value in document_text(classic_bytes)
        assert value in document_text(modern_bytes)
    body = DRAFT.splitlines()
    assert [t for t in texts_of(classic) if t in body] == [t for t in texts_of(modern) if t in body]
    assert classic.styles["Normal"].font.name == "Times New Roman"
    assert modern.styles["Normal"].font.name == "Arial"
    assert classic.styles["Normal"].font.size.pt == 14
    assert modern.styles["Normal"].font.size.pt == 12
    assert abs(classic.sections[0].left_margin - Mm(30)) < Mm(0.1)
    assert abs(modern.sections[0].top_margin - Mm(15)) < Mm(0.1)


def test_editor_formatting_reaches_the_downloaded_docx(client):
    result = prepared(client, requisites=filled("service_memo"))
    formatting = {
        "font": "Georgia",
        "font_size": 16,
        "line_spacing": 2,
        "paragraph_space_after_pt": 12,
        "first_line_indent_mm": 15,
        "body_alignment": "center",
        "body_bold": True,
        "body_italic": True,
        "body_underline": True,
    }
    rendered = Document(BytesIO(download(
        client, result["document"], "classic", formatting,
    ).content))
    body = [paragraph for paragraph in rendered.paragraphs if paragraph.text in DRAFT.splitlines()]

    assert rendered.styles["Normal"].font.name == "Georgia"
    assert rendered.styles["Normal"].font.size.pt == 16
    assert rendered.styles["Normal"].paragraph_format.line_spacing == 2
    assert rendered.styles["Normal"].paragraph_format.space_after.pt == 12
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER for paragraph in body)
    assert all(abs(paragraph.paragraph_format.first_line_indent - Mm(15)) < Mm(0.1) for paragraph in body)
    assert all(run.bold and run.italic and run.underline for paragraph in body for run in paragraph.runs)


@pytest.mark.parametrize("template_id", list(templates()))
def test_styles_use_template_font_without_theme_leftovers(client, template_id):
    doc = docx(client, template_id=template_id)
    font = templates()[template_id].font
    theme_names = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
    theme_attrs = [qn(f"w:{name}") for name in theme_names]
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault"))
    rpr_elements = [
        defaults.find(qn("w:rPr")),
        *(doc.styles[name].element.rPr for name in ("Normal", "Title", "Header", "Footer")),
    ]
    for rpr in rpr_elements:
        rfonts = rpr.find(qn("w:rFonts"))
        assert not any(attr in rfonts.attrib for attr in theme_attrs)
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            assert rfonts.get(qn(f"w:{name}")) == font
    assert defaults.find(qn("w:rPr")).find(qn("w:lang")).get(qn("w:val")) == "ru-RU"
    assert doc.styles["Title"].element.pPr.find(qn("w:pBdr")) is None


def test_text_reaches_the_file_character_for_character(client):
    draft = "ООО «Ромашка» — счёт № 45 от 03.11.2026, сумма 1 250,50 ₽; ёлка."
    result = prepared(client, "letter", {"subject": "Счёт № 15"}, draft=draft)
    text = document_text(download(client, result["document"]).content)
    assert draft in text
    assert "Счёт № 15" in text


def test_classic_template_places_requisites_as_in_the_reference(client):
    doc = docx(client, requisites=filled("service_memo"))
    texts = texts_of(doc)
    section = doc.sections[0]
    # Organization lives in the page header, the page number in the footer.
    assert texts_of(section.header) == [REQUISITES["organization"]]
    assert REQUISITES["organization"] not in texts
    assert 'w:instr="PAGE"' in section.footer._element.xml
    assert doc.styles["Header"].font.size.pt == 11
    # Requisites are printed without form labels such as «Кому:» or «Тема:».
    assert not any(text.startswith(("Кому:", "От кого:", "Тема:", "Дата:")) for text in texts)
    addressee = doc.paragraphs[texts.index(REQUISITES["recipient"])]
    assert addressee.paragraph_format.left_indent > Mm(50)
    signature = "Заведующий лабораторией\tИ. И. Иванов"
    order = [texts.index(value) for value in (
        REQUISITES["recipient"], REQUISITES["sender"], "СЛУЖЕБНАЯ ЗАПИСКА",
        "11.09.2026 № 47-СЗ", "О закупке мониторов", DRAFT.splitlines()[1],
        "Приложение: акт на 1 л. в 1 экз.", signature, REQUISITES["executor"],
    )]
    assert order == sorted(order)


def test_modern_template_places_requisites_as_in_the_reference(client):
    doc = docx(client, template_id="modern", requisites=filled("service_memo"))
    texts = texts_of(doc)
    section = doc.sections[0]
    assert not any(texts_of(section.header))
    assert texts_of(section.footer) == ["Служебная записка от 11.09.2026"]
    assert doc.styles["Footer"].font.size.pt == 10
    assert texts[0] == REQUISITES["organization"]
    rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert rows == [["Кому", REQUISITES["recipient"]], ["От кого", REQUISITES["sender"]]]
    for text in ("Заведующий лабораторией", "И. И. Иванов"):
        paragraph = doc.paragraphs[texts.index(text)]
        assert paragraph.alignment == 1  # center


def test_letter_shows_salutation_before_text_and_executor_at_the_end(client):
    texts = texts_of(docx(client, "letter", requisites=filled("letter")))
    order = [texts.index(value) for value in (
        "11.09.2026 № 47-СЗ", REQUISITES["recipient"], "ПИСЬМО", REQUISITES["subject"],
        REQUISITES["salutation"], DRAFT.splitlines()[0], REQUISITES["executor"],
    )]
    assert order == sorted(order)


def test_empty_optional_requisites_leave_no_traces(client):
    required = {field.id: REQUISITES[field.id]
                for field in document_types()["service_memo"].fields if field.required}
    for template_id in templates():
        document = prepared(client, requisites=required)["document"]
        text = document_text(download(client, document, template_id).content)
        assert "Приложение" not in text
        assert REQUISITES["executor"] not in text
        assert REQUISITES["organization"] not in text


def test_required_requisites_follow_the_case_reference():
    required = {
        type_id: {field.id for field in doc_type.fields if field.required}
        for type_id, doc_type in document_types().items()
    }
    assert required == {
        "service_memo": {"recipient", "sender", "date", "number", "subject", "signer"},
        "report_memo": {"recipient", "sender", "date", "number", "subject", "signer"},
        "information_note": {"subject", "date", "signer"},
        "letter": {"recipient", "organization", "date", "number", "subject", "signer"},
    }


def test_combined_requisites_must_be_adjacent():
    doc_type = DocumentType.model_validate({
        "id": "broken", "name": "x", "description": "x", "title": "X",
        "fields": [
            {"id": "date", "label": "Дата", "placement": "registration"},
            {"id": "number", "label": "Номер", "placement": "registration"},
        ],
        "blocks": ["date", "title", "number", "body"],
    })
    with pytest.raises(ValueError):
        check_contiguous_placements(doc_type, "broken.yaml")


@pytest.mark.parametrize("patch", [
    {"draft": "   "}, {"draft": "x" * 20001}, {"draft": "bad\x00text"},
    {"doc_type": "missing"}, {"requisites": {"invented": "value"}},
    {"requisites": {"recipient": "\x00"}}, {"extra": "not allowed"},
])
def test_invalid_processing_requests_are_rejected(client, patch):
    payload = {"draft": DRAFT, "doc_type": "service_memo", "requisites": {}} | patch
    assert client.post("/api/process", json=payload).status_code == 422


def test_download_validates_content_and_template(client):
    document = prepared(client)["document"]
    for payload in [
        {"document": document, "template_id": "unknown"},
        {"document": document | {"body": []}, "template_id": "classic"},
        {"document": document | {"body": ["bad\x00"]}, "template_id": "classic"},
        {"document": document | {"requisites": {"unknown": "x"}}, "template_id": "classic"},
        {"document": document, "template_id": "classic", "formatting": {
            "font": "Comic Sans MS", "font_size": 40, "line_spacing": 3,
            "paragraph_space_after_pt": 0, "first_line_indent_mm": 0,
            "body_alignment": "diagonal", "body_bold": False,
            "body_italic": False, "body_underline": False,
        }},
    ]:
        assert client.post("/api/documents/download", json=payload).status_code == 422


def test_custom_template_is_validated_and_used_for_download(client):
    document = prepared(client)["document"]
    custom = client.get("/api/catalog").json()["templates"][0] | {
        "id": "custom_editor_test",
        "name": "Моё оформление",
        "font": "Arial",
        "font_size": 13,
    }
    response = client.post("/api/documents/download", json={
        "document": document,
        "template_id": custom["id"],
        "custom_template": custom,
    })
    assert response.status_code == 200
    generated = Document(BytesIO(response.content))
    assert generated.styles["Normal"].font.name == "Arial"
    assert generated.styles["Normal"].font.size.pt == 13


def test_custom_template_id_must_match_selection(client):
    document = prepared(client)["document"]
    custom = client.get("/api/catalog").json()["templates"][0] | {
        "id": "custom_editor_test",
    }
    response = client.post("/api/documents/download", json={
        "document": document,
        "template_id": "custom_another",
        "custom_template": custom,
    })
    assert response.status_code == 422


def test_processor_failure_is_explicit_and_prepared_document_can_still_download(client):
    document = prepared(client)["document"]
    with TestClient(create_app(Settings(text_processor="unavailable"))) as failing:
        assert failing.get("/api/health").status_code == 200
        for _ in range(2):
            response = failing.post("/api/process", json={
                "draft": DRAFT, "doc_type": "service_memo", "requisites": {},
            })
            assert response.status_code == 503
            assert "недоступна" in response.json()["detail"]
        # Export accepts already prepared content and does not call the processor again.
        download(failing, document, "modern")
