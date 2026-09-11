"""Проверка готового DOCX: правила выбранного шаблона действительно попали в файл,
а состав и порядок реквизитов соответствуют выбранному типу документа."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from qa_support import (
    ALIGNMENT_NAMES,
    case_by_id,
    docx_text,
    download,
    measured_layout,
    process,
)

NEUTRAL_DRAFT = case_by_id("edge_01_single_line").draft


@pytest.fixture(scope="module")
def filled(client, catalog):
    """Документы всех типов со всеми заполненными реквизитами — без меток «Заполнить»."""
    documents = {}
    for doc_type in catalog["doc_types"]:
        values = {field["id"]: f"Значение {field['id']}" for field in doc_type["fields"]}
        documents[doc_type["id"]] = process(client, NEUTRAL_DRAFT, doc_type["id"], values)
    return documents


def paragraph_texts(data: bytes) -> list[str]:
    return [paragraph.text for paragraph in Document(BytesIO(data)).paragraphs]


def test_every_template_produces_a4_pages(client, filled, template_rules):
    for template_id in template_rules:
        data = download(client, filled["service_memo"]["document"], template_id)

        assert measured_layout(data)["page_mm"] == (210, 297), f"Шаблон «{template_id}»: не A4"


@pytest.mark.parametrize("type_id", ["service_memo", "report_memo", "information_note", "letter"])
def test_template_rules_are_written_into_the_file(client, filled, template_rules, type_id):
    for template_id, rules in template_rules.items():
        data = download(client, filled[type_id]["document"], template_id)

        layout = measured_layout(data)

        assert layout["font"] == rules["font"]
        assert layout["font_size"] == rules["font_size"]
        assert layout["line_spacing"] == rules["line_spacing"]
        assert layout["space_after_pt"] == rules["paragraph_space_after_pt"]
        for side, expected in rules["margins_mm"].items():
            assert abs(layout["margins_mm"][side] - expected) < 0.2, (
                f"Шаблон «{template_id}»: поле {side} не совпадает с правилом"
            )


def test_title_alignment_follows_the_template(client, filled, template_rules, catalog):
    titles = {doc_type["id"]: doc_type["title"] for doc_type in catalog["doc_types"]}
    for template_id, rules in template_rules.items():
        data = download(client, filled["letter"]["document"], template_id)

        document = Document(BytesIO(data))
        title = next(p for p in document.paragraphs if p.text == titles["letter"])

        assert ALIGNMENT_NAMES[title.alignment] == rules["title_alignment"]


def test_font_is_declared_for_latin_and_cyrillic_alike(client, filled, template_rules):
    for template_id, rules in template_rules.items():
        data = download(client, filled["service_memo"]["document"], template_id)

        fonts = Document(BytesIO(data)).styles["Normal"].element.rPr.rFonts

        # Word подбирает шрифт по алфавиту: без w:hAnsi кириллица может уехать на другой шрифт.
        assert fonts.get(qn("w:ascii")) == rules["font"], f"Шаблон «{template_id}»"
        assert fonts.get(qn("w:hAnsi")) == rules["font"], f"Шаблон «{template_id}»"


def test_requisites_appear_in_the_order_set_by_the_document_type(client, filled, catalog):
    for doc_type in catalog["doc_types"]:
        prepared = filled[doc_type["id"]]
        data = download(client, prepared["document"], "classic")
        texts = paragraph_texts(data)
        labels = {field["id"]: field["label"] for field in doc_type["fields"]}
        body = {line.strip() for line in prepared["document"]["body"]}

        positions = []
        for block in doc_type["blocks"]:
            if block == "title":
                positions.append(texts.index(doc_type["title"]))
            elif block == "body":
                positions.append(next(i for i, t in enumerate(texts) if t.strip() in body))
            else:
                prefix = f"{labels[block]}: "
                positions.append(next(i for i, t in enumerate(texts) if t.startswith(prefix)))

        assert positions == sorted(positions), (
            f"Тип «{doc_type['name']}»: порядок блоков в файле не совпадает с конфигурацией"
        )


def test_filled_requisites_leave_no_placeholders(client, filled, catalog):
    for doc_type in catalog["doc_types"]:
        data = download(client, filled[doc_type["id"]]["document"], "classic")

        text = docx_text(data)

        assert "[Заполнить:" not in text, f"Тип «{doc_type['name']}»: осталась метка пропуска"
        for field in doc_type["fields"]:
            assert f"Значение {field['id']}" in text, (
                f"Тип «{doc_type['name']}»: реквизит «{field['label']}» не попал в документ"
            )


def test_empty_optional_requisite_is_left_out_entirely(client, catalog):
    for doc_type in catalog["doc_types"]:
        optional = [field for field in doc_type["fields"] if not field["required"]]
        if not optional:
            continue
        values = {
            field["id"]: f"Значение {field['id']}"
            for field in doc_type["fields"] if field["required"]
        }
        prepared = process(client, NEUTRAL_DRAFT, doc_type["id"], values)

        text = docx_text(download(client, prepared["document"], "classic"))

        for field in optional:
            assert f"{field['label']}:" not in text, (
                f"Пустой необязательный реквизит «{field['label']}» занял место в документе"
            )


def test_typography_and_cyrillic_survive_the_file(client):
    draft = 'ООО «Ромашка» — счёт № 45 от 03.11.2026, сумма 1 250,50 ₽; ёлка, «кавычки».'

    prepared = process(client, draft, "letter")
    data = download(client, prepared["document"], "classic")

    assert draft in docx_text(data), "Кавычки, тире или буква «ё» изменились при записи в файл"


def test_file_carries_no_author_of_its_own(client, filled):
    data = download(client, filled["letter"]["document"], "classic")

    properties = Document(BytesIO(data)).core_properties

    assert not properties.author, "В свойствах файла остался автор"
    assert not properties.last_modified_by, "В свойствах файла остался редактор"
