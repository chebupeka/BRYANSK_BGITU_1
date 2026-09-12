"""Общие средства проверки: корпус черновиков, сравнение фактов и разбор готового DOCX.

Модуль намеренно не импортирует внутренние части обработки текста: проверки идут
через публичный API, поэтому они одинаково работают с заглушкой и с реальной моделью.
"""

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

CORPUS_DIR = Path(__file__).resolve().parents[2] / "examples" / "drafts"

# Пробел внутри числа — разделитель разрядов, а не граница значения: 180 000 == 180000.
DIGIT_GROUP = re.compile(r"(?<=\d)[\s ](?=\d)")
NUMBER = re.compile(r"\d[\d\s ]*\d|\d")
REVISION = re.compile(r"<w:(ins|del)[ />]")

ALIGNMENT_NAMES = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def normalized(text: str) -> str:
    """Приводит текст к виду, в котором сравнение не зависит от регистра и разрядов чисел."""
    return DIGIT_GROUP.sub("", text).casefold()


def keeps(haystack: str, value: str) -> bool:
    return normalized(value) in normalized(haystack)


def significant_numbers(text: str) -> set[str]:
    """Числа от двух цифр: суммы, количества, годы. Одиночные цифры дают слишком много шума."""
    found = {DIGIT_GROUP.sub("", match.group()) for match in NUMBER.finditer(text)}
    return {number for number in found if len(number) > 1}


@dataclass(frozen=True)
class Case:
    """Один вход корпуса вместе с ожидаемым структурным результатом."""

    id: str
    file: str
    doc_type: str
    label: str
    dirt: str
    must_keep: list[str]
    in_draft: dict[str, str]
    not_in_draft: list[str]

    @property
    def draft(self) -> str:
        return (CORPUS_DIR / self.file).read_text(encoding="utf-8")


def load_cases() -> list[Case]:
    manifest = yaml.safe_load((CORPUS_DIR / "manifest.yaml").read_text(encoding="utf-8"))
    return [Case(**case) for case in manifest["cases"]]


def case_by_id(case_id: str) -> Case:
    return next(case for case in load_cases() if case.id == case_id)


def process(client, draft: str, doc_type: str, requisites: dict[str, str] | None = None) -> dict:
    response = client.post("/api/process", json={
        "draft": draft, "doc_type": doc_type, "requisites": requisites or {},
    })
    assert response.status_code == 200, response.text
    return response.json()


def download(client, document: dict, template_id: str) -> bytes:
    response = client.post("/api/documents/download", json={
        "document": document, "template_id": template_id,
    })
    assert response.status_code == 200, response.text
    assert "wordprocessingml.document" in response.headers["content-type"]
    assert response.headers["content-disposition"].endswith('.docx"')
    return response.content


def prepared_text(document: dict) -> str:
    """Весь текст подготовленного документа: и реквизиты, и абзацы."""
    return "\n".join([*document["requisites"].values(), *document["body"]])


def docx_text(data: bytes) -> str:
    """Весь видимый текст файла: абзацы, ячейки таблиц и колонтитулы."""
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    for section in document.sections:
        for container in (section.header, section.footer):
            parts += [paragraph.text for paragraph in container.paragraphs]
    return "\n".join(parts)


def assert_editable_docx(data: bytes) -> None:
    """Файл должен открываться Word без восстановления и содержать правимый текст."""
    with ZipFile(BytesIO(data)) as archive:
        names = set(archive.namelist())
        assert archive.testzip() is None, "Архив DOCX повреждён"
        assert "[Content_Types].xml" in names, "Нет описания типов содержимого"
        assert "word/document.xml" in names, "Нет основной части документа"
        assert not any(name.startswith("word/media/") for name in names), (
            "Текст подменён картинкой — такой документ нельзя редактировать"
        )
        body = archive.read("word/document.xml").decode("utf-8")

    # Невалидный XML — та самая ошибка, после которой Word предлагает восстановить файл.
    ElementTree.fromstring(body)
    assert not REVISION.search(body), "В файле остались непринятые исправления"

    document = Document(BytesIO(data))
    assert any(paragraph.text.strip() for paragraph in document.paragraphs), "Документ пуст"
    assert "None" not in docx_text(data), "В текст попало служебное значение None"


def measured_layout(data: bytes) -> dict:
    """Оформление, реально записанное в файл, — его сверяют с правилами шаблона."""
    document = Document(BytesIO(data))
    section = document.sections[0]
    normal = document.styles["Normal"].paragraph_format
    font = document.styles["Normal"].font
    return {
        "font": font.name,
        "font_size": font.size.pt,
        "line_spacing": normal.line_spacing,
        "space_after_pt": normal.space_after.pt,
        "page_mm": (round(section.page_width.mm), round(section.page_height.mm)),
        "margins_mm": {
            "top": round(section.top_margin.mm, 1),
            "right": round(section.right_margin.mm, 1),
            "bottom": round(section.bottom_margin.mm, 1),
            "left": round(section.left_margin.mm, 1),
        },
    }


def body_paragraphs(data: bytes, body: list[str]) -> list:
    """Абзацы основного текста файла, найденные по тексту подготовленного документа."""
    wanted = {text.strip() for text in body if text.strip()}
    document = Document(BytesIO(data))
    return [
        paragraph for paragraph in document.paragraphs if paragraph.text.strip() in wanted
    ]
