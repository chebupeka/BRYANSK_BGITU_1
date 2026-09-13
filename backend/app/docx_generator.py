from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.schemas import DocumentContent, DocumentType, RequisiteField, Template

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
LANGUAGE = "ru-RU"
# Width of the label column in the «Кому / От кого» table.
TABLE_LABEL_WIDTH_MM = 35
# Word prefers theme font attributes over explicit names, so they must be removed.
THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
FONT_ATTRS = ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs")

# Requisites printed as compact blocks with single line spacing.
COMPACT = {"letterhead", "addressee", "registration", "headline", "signature", "executor"}
WIDE_GAP = {"signature", "executor"}


def set_fonts(rpr, font: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in THEME_FONT_ATTRS:
        rfonts.attrib.pop(qn(attr), None)
    for attr in FONT_ATTRS:
        rfonts.set(qn(attr), font)


def remove_children(parent, *tags: str) -> None:
    if parent is None:
        return
    for tag in tags:
        for child in parent.findall(qn(tag)):
            parent.remove(child)


def apply_document_defaults(doc, template: Template) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(rpr, template.font)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), LANGUAGE)


def apply_styles(doc, template: Template) -> None:
    apply_document_defaults(doc, template)
    normal = doc.styles["Normal"]
    set_fonts(normal.element.get_or_add_rPr(), template.font)
    normal.font.size = Pt(template.font_size)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = template.line_spacing
    normal.paragraph_format.space_after = Pt(template.paragraph_space_after_pt)
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title_rpr = title.element.get_or_add_rPr()
    set_fonts(title_rpr, template.font)
    # The default python-docx Title has a blue bottom border, letter spacing and kerning.
    remove_children(title.element.pPr, "w:pBdr")
    remove_children(title_rpr, "w:spacing", "w:kern", "w:szCs")
    title.font.size = Pt(template.font_size + 2)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    for name in ("Header", "Footer"):
        style = doc.styles[name]
        set_fonts(style.element.get_or_add_rPr(), template.font)
        style.font.size = Pt(template.header_footer_font_size)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    page_field.append(run)
    paragraph._p.append(page_field)


@dataclass
class Block:
    kind: str
    items: list[tuple[RequisiteField, str]] = field(default_factory=list)


def arrange_blocks(content: DocumentContent, doc_type: DocumentType) -> list[Block]:
    """Skip empty optional requisites and merge neighbours with the same placement."""
    fields = {item.id: item for item in doc_type.fields}
    blocks: list[Block] = []
    for name in doc_type.blocks:
        if name in {"title", "body"}:
            if name == "body" or doc_type.show_title:
                blocks.append(Block(name))
            continue
        requisite = fields[name]
        value = content.requisites.get(name, "").strip()
        if not value and not requisite.required:
            continue
        if not blocks or blocks[-1].kind != requisite.placement:
            blocks.append(Block(requisite.placement))
        blocks[-1].items.append((requisite, value))
    return blocks


def requisite_text(requisite: RequisiteField, value: str) -> str:
    return requisite.prefix + (value or f"[Заполнить: {requisite.label}]")


class Writer:
    def __init__(self, doc, template: Template, doc_type: DocumentType, content: DocumentContent):
        self.doc = doc
        self.template = template
        self.doc_type = doc_type
        self.content = content
        margins = template.margins_mm
        self.text_width_mm = PAGE_WIDTH_MM - margins["left"] - margins["right"]
        self.gap = Pt(template.font_size)

    def paragraph(self, text: str = "", alignment: str = "left", container=None):
        paragraph = (container or self.doc).add_paragraph(text)
        paragraph.alignment = ALIGNMENTS[alignment]
        return paragraph

    def write(self, block: Block) -> list | None:
        """Add a block to the page body. None means the block went to the page header."""
        template = self.template
        texts = [requisite_text(requisite, value) for requisite, value in block.items]
        if block.kind == "title":
            paragraph = self.doc.add_paragraph(self.doc_type.title, style="Title")
            paragraph.alignment = ALIGNMENTS[template.title_alignment]
            return [paragraph]
        if block.kind == "body":
            paragraphs = []
            for text in self.content.body:
                paragraph = self.paragraph(text, template.body_alignment)
                paragraph.paragraph_format.first_line_indent = Mm(template.first_line_indent_mm)
                for run in paragraph.runs:
                    run.bold = template.body_bold
                    run.italic = template.body_italic
                    run.underline = template.body_underline
                paragraphs.append(paragraph)
            return paragraphs
        if block.kind == "letterhead":
            if template.header == "organization":
                self.write_header(texts)
                return None
            return [self.paragraph(text, template.letterhead_alignment) for text in texts]
        if block.kind == "addressee":
            if template.recipient_layout == "table":
                self.addressee_table(block)
                return []
            return [self.addressee(text) for text in texts]
        if block.kind == "registration":
            return [self.paragraph(" ".join(texts))]
        if block.kind == "headline":
            paragraph = self.paragraph(" ".join(texts), template.headline_alignment)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = template.headline_bold
            return [paragraph]
        if block.kind == "salutation":
            return [self.paragraph(text, template.title_alignment) for text in texts]
        if block.kind == "signature":
            return self.signature(texts)
        if block.kind == "executor":
            paragraphs = [self.paragraph(text) for text in texts]
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(template.small_font_size)
            return paragraphs
        if block.kind == "paragraph":
            return [self.paragraph(text) for text in texts]
        return [self.labeled(requisite, value) for requisite, value in block.items]

    def write_header(self, texts: list[str]) -> None:
        header = self.doc.sections[0].header
        first = header.paragraphs[0]
        first.text = texts[0]
        first.alignment = ALIGNMENTS[self.template.letterhead_alignment]
        for text in texts[1:]:
            self.paragraph(text, self.template.letterhead_alignment, container=header)

    def write_footer(self) -> None:
        template = self.template
        if template.footer == "none":
            return
        paragraph = self.doc.sections[0].footer.paragraphs[0]
        if template.footer == "page_number":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_page_number(paragraph)
            return
        date = self.content.requisites.get("date", "").strip()
        paragraph.text = f"{self.doc_type.name} от {date}" if date else self.doc_type.name
        paragraph.alignment = ALIGNMENTS[template.letterhead_alignment]

    def addressee(self, text: str):
        alignment = self.template.recipient_alignment
        if alignment == "right":
            # Standard corner block: lines start in the right part of the page and wrap there.
            paragraph = self.paragraph(text)
            indent = self.text_width_mm - self.template.addressee_width_mm
            paragraph.paragraph_format.left_indent = Mm(indent)
            return paragraph
        return self.paragraph(text, alignment)

    def addressee_table(self, block: Block) -> None:
        """Two columns: «Кому | Руководителю…», «От кого | …»."""
        table = self.doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        widths = (Mm(TABLE_LABEL_WIDTH_MM), Mm(self.text_width_mm - TABLE_LABEL_WIDTH_MM))
        for column, width in zip(table.columns, widths, strict=True):
            column.width = width
        for requisite, value in block.items:
            cells = table.add_row().cells
            cells[0].paragraphs[0].add_run(requisite.label).bold = True
            cells[1].paragraphs[0].add_run(requisite_text(requisite, value))
            for cell, width in zip(cells, widths, strict=True):
                cell.width = width
                cell.paragraphs[0].paragraph_format.line_spacing = 1
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    def signature(self, texts: list[str]) -> list:
        alignment = self.template.signature_alignment
        if alignment == "left" and len(texts) > 1:
            # Position on the left, name at the right margin: «Заведующий   И. И. Иванов».
            paragraph = self.paragraph(f"{' '.join(texts[:-1])}\t{texts[-1]}")
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Mm(self.text_width_mm), WD_TAB_ALIGNMENT.RIGHT
            )
            return [paragraph]
        return [self.paragraph(text, alignment) for text in texts]

    def labeled(self, requisite: RequisiteField, value: str):
        paragraph = self.paragraph()
        paragraph.add_run(f"{requisite.label}: ").bold = True
        paragraph.add_run(value or f"[Заполнить: {requisite.label}]")
        return paragraph

    def space_before(self, kind: str, previous: str | None):
        if previous is None or previous == "title":
            return Pt(0)
        return self.gap * 2 if kind in WIDE_GAP else self.gap


def generate_docx(content: DocumentContent, doc_type: DocumentType, template: Template) -> bytes:
    """Type controls content and placement; template controls presentation, never wording."""
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = doc_type.title
    doc.core_properties.language = LANGUAGE
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(PAGE_WIDTH_MM), Mm(PAGE_HEIGHT_MM)
    for side, value in template.margins_mm.items():
        setattr(section, f"{side}_margin", Mm(value))
    # Header and footer sit in the middle of the margin, not glued to the text.
    section.header_distance = Mm(template.margins_mm["top"] / 2)
    section.footer_distance = Mm(template.margins_mm["bottom"] / 2)
    apply_styles(doc, template)

    writer = Writer(doc, template, doc_type, content)
    writer.write_footer()
    previous, previous_last = None, None
    for block in arrange_blocks(content, doc_type):
        paragraphs = writer.write(block)
        if paragraphs is None:
            continue
        if paragraphs and block.kind != "title":
            paragraphs[0].paragraph_format.space_before = writer.space_before(block.kind, previous)
        if not paragraphs and previous_last is not None:
            # A table cannot have a gap above it, so the paragraph before it gets one below.
            previous_last.paragraph_format.space_after = writer.gap
        if block.kind == "signature" and previous_last is not None:
            # The signature must not be alone on a page.
            previous_last.paragraph_format.keep_with_next = True
        if block.kind in COMPACT:
            # Addressee lines of different fields («кому» / «от кого») get a small gap.
            inner_gap = Pt(template.font_size / 2) if block.kind == "addressee" else Pt(0)
            for paragraph in paragraphs:
                paragraph.paragraph_format.line_spacing = 1
            for paragraph in paragraphs[:-1]:
                paragraph.paragraph_format.space_after = inner_gap
        previous = block.kind
        previous_last = paragraphs[-1] if paragraphs else None

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
